import io

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _style_header(cell, color="1F77B4"):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_subheader(cell):
    cell.font = Font(bold=True, size=10)
    cell.alignment = Alignment(wrap_text=True)


class ExcelReportGenerator:
    def generate_report(self, company_name: str, all_analysis: dict) -> bytes:
        if not OPENPYXL_AVAILABLE:
            return b""
        try:
            return self._build_excel(company_name, all_analysis)
        except Exception:
            return b""

    def _build_excel(self, company_name: str, all_analysis: dict) -> bytes:
        wb = openpyxl.Workbook()
        enrich = all_analysis.get("enrichment", {})
        financial = all_analysis.get("financial", {})
        market = all_analysis.get("market", {})
        competitive = all_analysis.get("competitive", {})
        strategic = all_analysis.get("strategic", {})
        valuation = all_analysis.get("valuation", {})

        # Sheet 1: Overview
        ws1 = wb.active
        ws1.title = "Overview"
        ws1.column_dimensions["A"].width = 30
        ws1.column_dimensions["B"].width = 50
        ws1.merge_cells("A1:B1")
        ws1["A1"] = f"Perfil Empresa: {company_name}"
        _style_header(ws1["A1"])
        ws1.row_dimensions[1].height = 25

        overview_data = [
            ("Nombre", company_name),
            ("Sector", enrich.get("sector", "N/D")),
            ("Subsector", enrich.get("subsector", "N/D")),
            ("Pais", enrich.get("country", "N/D")),
            ("Empleados", enrich.get("employees", "N/D")),
            ("Tipo empresa", enrich.get("company_type", "N/D")),
            ("Antiguedad (anos)", enrich.get("age_years", "N/D")),
            ("Tamano", enrich.get("size", "N/D")),
            ("Posicionamiento", enrich.get("market_positioning", "N/D")),
            ("Website", enrich.get("website", "N/D")),
        ]
        for i, (label, value) in enumerate(overview_data, start=2):
            ws1[f"A{i}"] = label
            ws1[f"B{i}"] = str(value) if value is not None else "N/D"
            _style_subheader(ws1[f"A{i}"])

        # Sheet 2: Financial
        ws2 = wb.create_sheet("Financial")
        ws2.column_dimensions["A"].width = 35
        ws2.column_dimensions["B"].width = 25
        ws2.merge_cells("A1:B1")
        ws2["A1"] = f"Analisis Financiero: {company_name}"
        _style_header(ws2["A1"])
        ws2.row_dimensions[1].height = 25

        fin_data = [
            ("Ingresos (M EUR)", financial.get("revenue", "N/D")),
            ("Beneficio Neto (M EUR)", financial.get("net_profit", "N/D")),
            ("EBITDA (M EUR)", financial.get("ebitda", "N/D")),
            ("Empleados", financial.get("employees", "N/D")),
            ("Crecimiento Anual (%)", financial.get("growth_rate", "N/D")),
            ("Margen Neto (%)", financial.get("margin", "N/D")),
            ("Ratio Deuda", financial.get("debt_ratio", "N/D")),
            ("Anio datos", financial.get("year", "N/D")),
            ("Fiabilidad", financial.get("reliability", "N/D")),
            ("Es estimacion", financial.get("is_estimated", "N/D")),
            ("Fuente", financial.get("data_source", "N/D")),
            ("Notas", financial.get("notes", "N/D")),
        ]
        for i, (label, value) in enumerate(fin_data, start=2):
            ws2[f"A{i}"] = label
            ws2[f"B{i}"] = str(value) if value is not None else "N/D"
            _style_subheader(ws2[f"A{i}"])

        trend = financial.get("revenue_trend_3y", [])
        if trend and len(trend) == 3:
            year = financial.get("year", 2024)
            ws2["A15"] = "Evolucion Ingresos 3 Anos"
            _style_header(ws2["A15"], "2E86AB")
            ws2["A16"] = str(year - 2)
            ws2["B16"] = trend[0]
            ws2["A17"] = str(year - 1)
            ws2["B17"] = trend[1]
            ws2["A18"] = str(year)
            ws2["B18"] = trend[2]

        # Sheet 3: SWOT
        ws3 = wb.create_sheet("SWOT")
        ws3.column_dimensions["A"].width = 25
        ws3.column_dimensions["B"].width = 60
        ws3.merge_cells("A1:B1")
        ws3["A1"] = f"Analisis SWOT: {company_name}"
        _style_header(ws3["A1"])
        ws3.row_dimensions[1].height = 25

        swot = strategic.get("swot", {})
        row = 2
        for category, label, color in [
            ("strengths", "FORTALEZAS", "70AD47"),
            ("weaknesses", "DEBILIDADES", "FF0000"),
            ("opportunities", "OPORTUNIDADES", "00B0F0"),
            ("threats", "AMENAZAS", "FFC000"),
        ]:
            ws3[f"A{row}"] = label
            _style_header(ws3[f"A{row}"], color)
            row += 1
            for item in swot.get(category, []):
                ws3[f"B{row}"] = f"- {item}"
                ws3[f"B{row}"].alignment = Alignment(wrap_text=True)
                row += 1
            row += 1

        # Sheet 4: Valuation
        ws4 = wb.create_sheet("Valuation")
        ws4.column_dimensions["A"].width = 35
        ws4.column_dimensions["B"].width = 30
        ws4.merge_cells("A1:B1")
        ws4["A1"] = f"Valoracion y Propension a Venta: {company_name}"
        _style_header(ws4["A1"])
        ws4.row_dimensions[1].height = 25

        val_obj = valuation.get("valuation", {})
        sale = valuation.get("sale_propensity", {})
        val_data = [
            ("Valor Estimado (M EUR)", val_obj.get("estimated_value_m", "N/D")),
            ("Valor Minimo (M EUR)", val_obj.get("min_value_m", "N/D")),
            ("Valor Maximo (M EUR)", val_obj.get("max_value_m", "N/D")),
            ("Confianza Valoracion", val_obj.get("confidence", "N/D")),
            ("Metodologia", val_obj.get("method", "N/D")),
            ("Multiplo EBITDA", val_obj.get("multiples_used", {}).get("ebitda_multiple", "N/D")),
            ("", ""),
            ("PROPENSION A LA VENTA", ""),
            ("Probabilidad Venta", sale.get("probability", "N/D")),
            ("Recomendacion", sale.get("recommendation", "N/D")),
            ("Compradores Potenciales", ", ".join(sale.get("buyer_types", []))),
        ]
        for i, (label, value) in enumerate(val_data, start=2):
            ws4[f"A{i}"] = label
            ws4[f"B{i}"] = str(value) if value is not None else "N/D"
            if label:
                _style_subheader(ws4[f"A{i}"])

        # Sheet 5: Competitors
        ws5 = wb.create_sheet("Competitors")
        headers = ["Competidor", "Cuota Mercado", "Fortalezas", "Debilidades"]
        for col, h in enumerate(headers, start=1):
            cell = ws5[f"{get_column_letter(col)}1"]
            cell.value = h
            _style_header(cell)
        ws5.column_dimensions["A"].width = 25
        ws5.column_dimensions["B"].width = 20
        ws5.column_dimensions["C"].width = 45
        ws5.column_dimensions["D"].width = 45

        for i, comp in enumerate(competitive.get("competitors", []), start=2):
            ws5[f"A{i}"] = comp.get("name", "N/D")
            ws5[f"B{i}"] = comp.get("market_share", "N/D")
            ws5[f"C{i}"] = ", ".join(comp.get("strengths", []))
            ws5[f"D{i}"] = ", ".join(comp.get("weaknesses", []))
            for col in range(1, 5):
                ws5[f"{get_column_letter(col)}{i}"].alignment = Alignment(wrap_text=True)

        # Sheet 6: Scenarios
        ws6 = wb.create_sheet("Scenarios")
        ws6.column_dimensions["A"].width = 25
        ws6.column_dimensions["B"].width = 70
        ws6.merge_cells("A1:B1")
        ws6["A1"] = f"Escenarios Futuros: {company_name}"
        _style_header(ws6["A1"])
        ws6.row_dimensions[1].height = 25

        scenario_data = [
            ("HORIZONTE 5 ANOS", ""),
            ("Escenario Optimista", strategic.get("scenario_5y", {}).get("best", "N/D")),
            ("Escenario Pesimista", strategic.get("scenario_5y", {}).get("worst", "N/D")),
            ("Escenario Mas Probable", strategic.get("scenario_5y", {}).get("most_likely", "N/D")),
            ("", ""),
            ("HORIZONTE 10 ANOS", ""),
            ("Escenario Optimista", strategic.get("scenario_10y", {}).get("best", "N/D")),
            ("Escenario Pesimista", strategic.get("scenario_10y", {}).get("worst", "N/D")),
            ("Escenario Mas Probable", strategic.get("scenario_10y", {}).get("most_likely", "N/D")),
        ]
        for i, (label, value) in enumerate(scenario_data, start=2):
            ws6[f"A{i}"] = label
            ws6[f"B{i}"] = str(value)
            if label:
                _style_subheader(ws6[f"A{i}"])
            ws6[f"B{i}"].alignment = Alignment(wrap_text=True)
            ws6.row_dimensions[i].height = 40

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
